import io
import pypdf
import docx


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """
    Extract text content from a PDF file using pypdf.

    Raises:
        ValueError: If no text could be extracted (e.g. scanned/image-only PDF).
    """
    try:
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        extracted_pages = []
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                extracted_pages.append(page_text)
        
        full_text = "\n".join(extracted_pages).strip()
        if not full_text:
            raise ValueError(
                "No readable text could be found in this PDF. "
                "It may be a scanned image or empty document. Please upload a text-based PDF."
            )
        return full_text
    except ValueError:
        raise
    except Exception as e:
        raise ValueError(
            f"Failed to parse PDF file. The file may be corrupt or encrypted. Details: {e}"
        )


def extract_text_from_docx(file_bytes: bytes) -> str:
    """
    Extract text content from a Word (.docx) document using python-docx.

    Raises:
        ValueError: If no text could be extracted from the document.
    """
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        extracted_text = []

        # Extract text from paragraphs
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                extracted_text.append(paragraph.text.strip())

        # Extract text from tables if any
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    extracted_text.append(" | ".join(row_text))

        full_text = "\n".join(extracted_text).strip()
        if not full_text:
            raise ValueError(
                "No readable text could be found in this Word document. "
                "Please ensure the document contains text content."
            )
        return full_text
    except ValueError:
        raise
    except Exception as e:
        raise ValueError(
            f"Failed to parse Word document. The file may be corrupt. Details: {e}"
        )


# Minimum word count before we warn the user about sparse extraction
_MIN_WORD_COUNT = 50


def extract_resume_text(uploaded_file) -> tuple[str, int]:
    """
    Determine file type from uploaded_file and invoke the appropriate parser.

    Args:
        uploaded_file: Streamlit UploadedFile object or file-like object.

    Returns:
        tuple[str, int]: (extracted_text, word_count).
        Callers should warn the user when word_count < _MIN_WORD_COUNT.

    Raises:
        ValueError: If the file is None, format is unsupported, or text
                    extraction fails (empty file, corrupt PDF, etc.).
    """
    if uploaded_file is None:
        raise ValueError("No file uploaded. Please upload a PDF or Word document.")

    filename = getattr(uploaded_file, "name", "").lower()

    if hasattr(uploaded_file, "getvalue"):
        file_bytes = uploaded_file.getvalue()
    elif hasattr(uploaded_file, "read"):
        file_bytes = uploaded_file.read()
    else:
        file_bytes = bytes(uploaded_file)

    if not file_bytes:
        raise ValueError(
            "The uploaded file appears to be empty (0 bytes). "
            "Please check the file and try again."
        )

    if filename.endswith(".pdf"):
        text = extract_text_from_pdf(file_bytes)
    elif filename.endswith(".docx") or filename.endswith(".doc"):
        text = extract_text_from_docx(file_bytes)
    else:
        raise ValueError(
            f"Unsupported file format '{filename}'. "
            "Please upload a PDF (.pdf) or Word document (.docx)."
        )

    word_count = len(text.split())
    return text, word_count
